import io
import tempfile
import zipfile
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.storage import FileSystemStorage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook
from PIL import Image

from catalog.bulk_import import BulkProductImportService
from catalog.models import Category, ImportJob, Product, ProductImage, ProductVariant


def workbook_bytes(rows, headers=None):
    headers = headers or [
        'product_code', 'name', 'description', 'tagline', 'category',
        'status', 'color', 'color_hex', 'size', 'sku', 'price', 'stock',
        'image_1', 'image_2', 'image_3', 'image_4',
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Products'
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, '') for header in headers])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def image_bytes(fmt='JPEG', color=(10, 20, 30)):
    output = io.BytesIO()
    Image.new('RGB', (20, 20), color).save(output, format=fmt)
    return output.getvalue()


def package_bytes(rows, images=None, include_images=True, excel_name='products.xlsx'):
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(excel_name, workbook_bytes(rows))
        if include_images:
            archive.writestr('images/luna.jpg', (images or {}).get('luna.jpg', image_bytes()))
    return output.getvalue()


class BulkImportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username='importer', password='password', is_staff=True)
        self.category = Category.objects.create(name='Dresses', slug='dresses')

    def run_import(self, rows, images=None, **kwargs):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        storage = FileSystemStorage(location=temporary.name)
        upload = io.BytesIO(package_bytes(rows, images=images, **kwargs))
        from django.core.files.uploadedfile import SimpleUploadedFile
        package = SimpleUploadedFile(
            'MODEZA_IMPORT.zip', upload.getvalue(), content_type='application/zip')
        with patch('catalog.bulk_import.default_storage', storage):
            job = BulkProductImportService.create_job(
                package, self.user, import_status='DRAFT')
            job = BulkProductImportService.validate(job, actor=self.user)
            return job, storage

    def test_valid_package_is_previewed_before_catalog_changes(self):
        rows = [
            {'product_code': 'LSD001', 'name': 'Luna Silk Dress', 'description': 'Silk dress', 'category': 'Dresses', 'color': 'Black', 'color_hex': '#000000', 'size': 'S', 'sku': 'LSD-BLK-S', 'price': '4500', 'stock': '5', 'image_1': 'luna.jpg'},
            {'product_code': 'LSD001', 'name': 'Luna Silk Dress', 'description': 'Silk dress', 'category': 'Dresses', 'color': 'Black', 'color_hex': '#000000', 'size': 'M', 'sku': 'LSD-BLK-M', 'price': '4500', 'stock': '7', 'image_1': 'luna.jpg'},
        ]
        job, _storage = self.run_import(rows)
        self.assertEqual(job.status, ImportJob.Status.READY)
        self.assertEqual(job.error_count, 0)
        self.assertEqual(job.validation_results['summary']['products'], 1)
        self.assertEqual(job.validation_results['summary']['variants'], 2)
        self.assertEqual(Product.objects.count(), 0)
        self.assertEqual(ProductVariant.objects.count(), 0)

    def test_confirm_and_process_creates_product_variants_and_inventory(self):
        rows = [
            {'product_code': 'LSD001', 'name': 'Luna Silk Dress', 'description': 'Silk dress', 'category': 'Dresses', 'color': 'Black', 'size': 'S', 'sku': 'LSD-BLK-S', 'price': '4500', 'stock': '5', 'image_1': 'luna.jpg'},
            {'product_code': 'LSD001', 'name': 'Luna Silk Dress', 'description': 'Silk dress', 'category': 'Dresses', 'color': 'Black', 'size': 'M', 'sku': 'LSD-BLK-M', 'price': '4500', 'stock': '7', 'image_1': 'luna.jpg'},
        ]
        job, storage = self.run_import(rows)
        with patch('catalog.bulk_import.default_storage', storage):
            job = BulkProductImportService.confirm(job, actor=self.user)
            job, finished = BulkProductImportService.process_chunk(job, limit=10, actor=self.user)
        self.assertTrue(finished)
        self.assertEqual(job.status, ImportJob.Status.COMPLETED)
        product = Product.objects.get(product_code='LSD001')
        self.assertEqual(product.variants.count(), 2)
        self.assertEqual(product.stock_quantity, 12)
        self.assertEqual(product.images, [storage.url('products/lsd001/luna.jpg')])
        self.assertEqual(ProductVariant.objects.get(sku='LSD-BLK-S').stock_quantity, 5)
        self.assertEqual(job.created_products, 1)
        self.assertEqual(job.created_variants, 2)
        self.assertEqual(job.uploaded_images, 1)

    def test_invalid_workbook_does_not_modify_catalog(self):
        rows = [{'product_code': 'LSD001', 'name': 'Luna', 'category': 'Unknown', 'color': 'Black', 'size': 'S', 'sku': 'LSD-S', 'price': 'bad', 'stock': '-1'}]
        job, _storage = self.run_import(rows)
        self.assertEqual(job.status, ImportJob.Status.READY)
        self.assertGreater(job.error_count, 0)
        self.assertFalse(job.validation_results['can_confirm'])
        self.assertEqual(Product.objects.count(), 0)

    def test_duplicate_sku_and_variant_are_errors(self):
        rows = [
            {'product_code': 'LSD001', 'name': 'Luna', 'category': 'Dresses', 'color': 'Black', 'size': 'S', 'sku': 'DUPLICATE', 'price': '10', 'stock': '1'},
            {'product_code': 'LSD001', 'name': 'Luna', 'category': 'Dresses', 'color': 'Black', 'size': 'S', 'sku': 'DUPLICATE', 'price': '10', 'stock': '1'},
        ]
        job, _storage = self.run_import(rows)
        messages = ' '.join(item['message'] for item in job.validation_results['errors'])
        self.assertIn('Duplicate SKU', messages)
        self.assertIn('Duplicate size and color', messages)

    def test_missing_images_directory_is_rejected(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as archive:
            archive.writestr('products.xlsx', workbook_bytes([
                {'product_code': 'LSD001', 'name': 'Luna', 'category': 'Dresses', 'color': 'Black', 'size': 'S', 'sku': 'LSD-S', 'price': '10', 'stock': '1'},
            ]))
        package = SimpleUploadedFile('missing-images.zip', output.getvalue(), content_type='application/zip')
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        storage = FileSystemStorage(location=temporary.name)
        with patch('catalog.bulk_import.default_storage', storage):
            job = BulkProductImportService.create_job(package, self.user)
            job = BulkProductImportService.validate(job, actor=self.user)
        self.assertEqual(job.status, ImportJob.Status.FAILED)
        self.assertIn('images/', ' '.join(item['message'] for item in job.validation_results['errors']))

    def test_corrupt_image_is_rejected_before_import(self):
        rows = [{'product_code': 'LSD001', 'name': 'Luna', 'category': 'Dresses', 'color': 'Black', 'size': 'S', 'sku': 'LSD-S', 'price': '10', 'stock': '1', 'image_1': 'luna.jpg'}]
        job, _storage = self.run_import(rows, images={'luna.jpg': b'not-an-image'})
        self.assertTrue(job.error_count)
        self.assertIn('cannot be decoded', ' '.join(item['message'] for item in job.validation_results['errors']))
        self.assertEqual(Product.objects.count(), 0)

    def test_missing_required_column_is_rejected(self):
        output = io.BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Products'
        sheet.append(['product_code', 'name', 'category', 'color', 'size', 'sku', 'price'])
        sheet.append(['LSD001', 'Luna', 'Dresses', 'Black', 'S', 'LSD-S', '10'])
        workbook_data = io.BytesIO()
        workbook.save(workbook_data)
        with zipfile.ZipFile(output, 'w') as archive:
            archive.writestr('products.xlsx', workbook_data.getvalue())
            archive.writestr('images/luna.jpg', image_bytes())
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        storage = FileSystemStorage(location=temporary.name)
        package = SimpleUploadedFile('missing-column.zip', output.getvalue(), content_type='application/zip')
        with patch('catalog.bulk_import.default_storage', storage):
            job = BulkProductImportService.create_job(package, self.user)
            job = BulkProductImportService.validate(job, actor=self.user)
        self.assertEqual(job.status, ImportJob.Status.FAILED)
        self.assertIn('Missing required columns: stock', ' '.join(item['message'] for item in job.validation_results['errors']))

    def test_missing_excel_and_unsafe_path_are_rejected(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as archive:
            archive.writestr('../products.xlsx', workbook_bytes([]))
            archive.writestr('images/luna.jpg', image_bytes())
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        storage = FileSystemStorage(location=temporary.name)
        package = SimpleUploadedFile('bad.zip', output.getvalue(), content_type='application/zip')
        with patch('catalog.bulk_import.default_storage', storage):
            job = BulkProductImportService.create_job(package, self.user)
            job = BulkProductImportService.validate(job, actor=self.user)
        self.assertEqual(job.status, ImportJob.Status.FAILED)
        self.assertTrue(job.error_count)
        self.assertEqual(Product.objects.count(), 0)

    def test_product_atomicity_rolls_back_when_image_upload_fails(self):
        rows = [{'product_code': 'LSD001', 'name': 'Luna Silk Dress', 'description': 'Silk dress', 'category': 'Dresses', 'color': 'Black', 'size': 'S', 'sku': 'LSD-BLK-S', 'price': '4500', 'stock': '5', 'image_1': 'luna.jpg'}]
        job, storage = self.run_import(rows)
        with patch('catalog.bulk_import.default_storage', storage), patch.object(storage, 'save', side_effect=OSError('storage unavailable')):
            BulkProductImportService.confirm(job, actor=self.user)
            job, finished = BulkProductImportService.process_chunk(job, limit=10, actor=self.user)
        self.assertTrue(finished)
        self.assertEqual(job.status, ImportJob.Status.COMPLETED_WITH_ERRORS)
        self.assertEqual(job.failed_rows, 1)
        self.assertEqual(Product.objects.count(), 0)
        self.assertEqual(ProductVariant.objects.count(), 0)

    def test_repeated_import_updates_without_duplicate_variants_or_images(self):
        rows = [{'product_code': 'LSD001', 'name': 'Luna Silk Dress', 'description': 'Silk dress', 'category': 'Dresses', 'color': 'Black', 'size': 'S', 'sku': 'LSD-BLK-S', 'price': '4500', 'stock': '5', 'image_1': 'luna.jpg'}]
        first_job, storage = self.run_import(rows)
        with patch('catalog.bulk_import.default_storage', storage):
            BulkProductImportService.confirm(first_job, actor=self.user)
            BulkProductImportService.process_chunk(first_job, limit=10, actor=self.user)
            second = BulkProductImportService.create_job(
                SimpleUploadedFile('MODEZA_IMPORT.zip', package_bytes(rows), content_type='application/zip'),
                self.user,
                import_status='DRAFT',
            )
            second = BulkProductImportService.validate(second, actor=self.user)
            BulkProductImportService.confirm(second, actor=self.user)
            second, finished = BulkProductImportService.process_chunk(second, limit=10, actor=self.user)
        self.assertTrue(finished)
        self.assertEqual(Product.objects.count(), 1)
        self.assertEqual(ProductVariant.objects.count(), 1)
        self.assertEqual(ProductImage.objects.count(), 1)
        self.assertEqual(second.updated_products, 1)
        self.assertEqual(second.created_variants, 0)
        self.assertEqual(second.updated_variants, 1)

    def test_template_contains_supported_columns(self):
        from catalog.services import BulkImportTemplateService
        workbook = BulkImportTemplateService.generate_workbook()
        headers = [cell.value for cell in workbook['Products'][1]]
        self.assertEqual(headers[:5], ['product_code', 'name', 'description', 'tagline', 'category'])
        self.assertNotIn('brand', headers)
        self.assertIn('Instructions', workbook.sheetnames)
        self.assertIn('Example', workbook.sheetnames)
